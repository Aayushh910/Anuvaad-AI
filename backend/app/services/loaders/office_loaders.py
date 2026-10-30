import io
import logging
from pathlib import Path
from app.services.loaders.base import BaseLoader, Segment, compute_file_hash
from app.services.loaders.image_processor import image_processor

logger = logging.getLogger(__name__)


class DocxLoader(BaseLoader):
    """Loader for Microsoft Word (.docx) documents."""

    def load(self, path: Path | str) -> list[Segment]:
        docx_path = Path(path).resolve()
        file_hash = compute_file_hash(docx_path)
        segments: list[Segment] = []

        try:
            import docx
            doc = docx.Document(str(docx_path))
        except Exception as e:
            logger.error("Failed to load DOCX %s: %s", docx_path.name, e)
            return []

        text_lines = []
        for p in doc.paragraphs:
            if p.text.strip():
                text_lines.append(p.text.strip())

        # Extract tables
        table_lines = []
        for t_idx, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                table_lines.append(f"[Table {t_idx}]:\n" + "\n".join(rows))

        # Extract embedded images
        image_parts = []
        try:
            img_count = 0
            for rel in doc.part.related_parts.values():
                if "image" in rel.content_type:
                    img_count += 1
                    img_bytes = rel.blob
                    ext = rel.content_type.split("/")[-1]
                    desc = image_processor.process_image(
                        image_bytes=img_bytes,
                        file_hash=file_hash,
                        index_or_page=f"docx_img_{img_count}",
                        ext=ext,
                    )
                    if desc:
                        image_parts.append(f"[Embedded image {img_count}]:\n{desc}")
        except Exception as e:
            logger.warning("Error reading images from DOCX %s: %s", docx_path.name, e)

        full_parts = []
        if text_lines:
            full_parts.append("\n\n".join(text_lines))
        if table_lines:
            full_parts.append("\n\n".join(table_lines))
        if image_parts:
            full_parts.append("\n\n".join(image_parts))

        combined = "\n\n".join(full_parts).strip()
        if combined:
            segments.append(
                Segment(
                    text=combined,
                    source=docx_path.name,
                    page_or_slide_or_sheet=1,
                    content_type="text",
                    metadata={"file_hash": file_hash},
                )
            )

        return segments


class PptxLoader(BaseLoader):
    """Loader for Microsoft PowerPoint (.pptx) presentations."""

    def load(self, path: Path | str) -> list[Segment]:
        pptx_path = Path(path).resolve()
        file_hash = compute_file_hash(pptx_path)
        segments: list[Segment] = []

        try:
            import pptx
            prs = pptx.Presentation(str(pptx_path))
        except Exception as e:
            logger.error("Failed to load PPTX %s: %s", pptx_path.name, e)
            return []

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_texts = []
            slide_tables = []
            slide_images = []

            for shape in slide.shapes:
                # Text frame
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        txt = para.text.strip()
                        if txt:
                            slide_texts.append(txt)

                # Tables
                elif shape.has_table:
                    table = shape.table
                    rows = []
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        rows.append(" | ".join(cells))
                    if rows:
                        slide_tables.append("\n".join(rows))

                # Slide Images / Pictures
                elif shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                    try:
                        img_bytes = shape.image.blob
                        ext = shape.image.ext
                        desc = image_processor.process_image(
                            image_bytes=img_bytes,
                            file_hash=file_hash,
                            index_or_page=f"slide{slide_idx}_{shape.shape_id}",
                            ext=ext,
                        )
                        if desc:
                            slide_images.append(f"[Slide image]:\n{desc}")
                    except Exception as e:
                        logger.warning("Error reading image on slide %s: %s", slide_idx, e)

            # Speaker Notes
            notes_text = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()

            parts = []
            if slide_texts:
                parts.append("\n".join(slide_texts))
            if slide_tables:
                parts.append("[Table]:\n" + "\n\n".join(slide_tables))
            if slide_images:
                parts.extend(slide_images)
            if notes_text:
                parts.append(f"[Speaker Notes]:\n{notes_text}")

            combined = "\n\n".join(parts).strip()
            if combined:
                segments.append(
                    Segment(
                        text=combined,
                        source=pptx_path.name,
                        page_or_slide_or_sheet=slide_idx,
                        content_type="text",
                        metadata={
                            "file_hash": file_hash,
                            "slide_number": slide_idx,
                            "has_visual_content": len(slide_images) > 0,
                        },
                    )
                )

        return segments

